#!/usr/bin/env python3
"""
html2docx_advanced.py — 高级层处理（10 种 IR）

IR 类型：quote / columns / compare_cards / comparison / faq / badge_group /
         gallery / pricing_table / icon_list / end_page

被 html2docx.py 主入口 import。
依赖 docx_utils + sem_common，跨层依赖 html2docx_core（process_table / _extract_img_bytes）。
"""

import io
import re

from bs4 import Tag, NavigableString, Comment

from sem_common import (
    get_text_content, find_semantic_class, parse_table_data, get_data_cols,
)
from docx_utils import (
    hex_to_rgb, is_dark_color,
    extract_bg_with_gradient_fallback, extract_bg_from_styles,
    extract_color_from_styles,
    set_run_font, add_cell_paragraph, flush_pending_page_break, inject_page_break_before,
    set_cell_margins, set_table_borders_nil, set_cell_borders, set_table_width_percent,
    apply_shading, apply_shading_to_cell, apply_shading_to_run,
)
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 跨层依赖：comparison 兜底调 process_table；gallery 调 _extract_img_bytes；
# columns/comparison/pricing_table 调 _add_rich_text_runs
from html2docx_core import process_table, _extract_img_bytes, _add_rich_text_runs


def process_quote(doc, section, palette):
    """引用语录 → 缩进段落（斜体 + 左边框 + 背景色 shading），多段落完整输出，保留 <strong>/<span> 等富文本"""
    bg_color = extract_bg_with_gradient_fallback(section.get("style", ""))
    ps = section.find_all("p")
    if not ps:
        # 尝试 blockquote 内部文本
        text = get_text_content(section)
        if not text:
            flush_pending_page_break(doc)
            return
        ps = [section]
    primary_color = palette.get("primary", "1e293b")
    for idx, p in enumerate(ps):
        if not get_text_content(p):
            continue
        para = doc.add_paragraph()
        if idx == 0 and flush_pending_page_break(doc):
            inject_page_break_before(para)
        if bg_color:
            apply_shading(para, bg_color)
        if idx == 0:
            # 首段：缩进 + 斜体 + 左边框 + 富文本写入
            _add_rich_text_runs(para, p, default_size=Pt(11), default_color=primary_color)
            for run in para.runs:
                run.italic = True
            para.paragraph_format.left_indent = Cm(1)
            pPr = para.paragraph_format.element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left_bdr = OxmlElement("w:left")
            left_bdr.set(qn("w:val"), "single")
            left_bdr.set(qn("w:sz"), "12")
            left_bdr.set(qn("w:space"), "8")
            left_bdr.set(qn("w:color"), palette.get("accent", "6366f1"))
            pBdr.append(left_bdr)
            pPr.append(pBdr)
        else:
            # 后续段落：右对齐 + 小字号浅色（如来源/署名）+ 富文本写入
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _add_rich_text_runs(para, p, default_size=Pt(10), default_color="64748b")
            if bg_color:
                apply_shading(para, bg_color)


def process_columns(doc, section, palette=None):
    """多栏并排 → 无边框表格列 + 卡片背景色 shading + 富文本写入"""
    cols = []
    for child in section.children:
        if isinstance(child, Tag) and child.get_text(strip=True):
            cols.append(child)
    if not cols:
        flush_pending_page_break(doc)
        return
    table = doc.add_table(rows=1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_nil(table)
    if flush_pending_page_break(doc):
        first_cell = table.cell(0, 0)
        if first_cell.paragraphs:
            inject_page_break_before(first_cell.paragraphs[0])
    text_color = (palette or {}).get("text", "1f2d3d")
    for j, col in enumerate(cols):
        cell = table.cell(0, j)
        bg = extract_bg_with_gradient_fallback(col.get("style", ""))
        if bg:
            apply_shading_to_cell(cell, bg)
        # 富文本写入：遍历列内所有子元素（p/h3/ul 等），保留格式
        first_para = True
        for child in col.children:
            if not isinstance(child, Tag):
                continue
            if child.name == "p":
                p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                first_para = False
                _add_rich_text_runs(p_cell, child, default_size=Pt(11), default_color=text_color)
            elif child.name in ("h3", "h4"):
                p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                first_para = False
                p_cell.paragraph_format.space_before = Pt(4)
                p_cell.paragraph_format.space_after = Pt(2)
                run = p_cell.add_run(get_text_content(child))
                set_run_font(run, size=Pt(13), bold=True, color=hex_to_rgb((palette or {}).get("primary", "1e293b")))
            elif child.name in ("ul", "ol"):
                for li in child.find_all("li", recursive=False):
                    p_cell = add_cell_paragraph(cell)
                    _add_rich_text_runs(p_cell, li, default_size=Pt(11), default_color=text_color)
            else:
                text = get_text_content(child)
                if text:
                    p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                    first_para = False
                    _add_rich_text_runs(p_cell, child, default_size=Pt(11), default_color=text_color)
        # 如果没有写入任何内容，兜底用纯文本
        if first_para:
            cell.text = col.get_text(strip=True)
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=Pt(11), color=hex_to_rgb(text_color))
        set_cell_margins(cell, top=30, start=60, bottom=30, end=60)


def process_compare_cards(doc, section, palette, parsed_styles=None, soup=None):
    """并排对比卡片（compare-grid/compare-card）→ 多列无边框表格

    每张卡片占一列：标签行（tagline）+ 标题（h3）+ 参数行 + 正文行。
    样式保留：白卡 → 淡边框 + 顶部 accent 粗线；推荐卡 → 渐变中间色底 + 白字；
    文字颜色读取 HTML 中 .compare-card 相关选择器的 color 定义，未定义时回退调色板。
    """
    cards = [c for c in section.children if isinstance(c, Tag)
             and "compare-card" in (c.get("class") or [])]
    if not cards:
        flush_pending_page_break(doc)
        return
    max_cols = get_data_cols(section)
    if max_cols and max_cols >= 2:
        cards = cards[:max_cols]
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_nil(table)
    set_table_width_percent(table, 100)
    if flush_pending_page_break(doc):
        first_cell = table.cell(0, 0)
        if first_cell.paragraphs:
            inject_page_break_before(first_cell.paragraphs[0])
    accent_color = palette.get("accent", "c9991f")
    border_color = palette.get("border", "e8dcbf")

    for j, card in enumerate(cards):
        cell = table.cell(0, j)
        classes = card.get("class") or []
        is_recommend = "recommend" in classes
        bg = extract_bg_with_gradient_fallback(card.get("style", ""))
        if not bg:
            if is_recommend:
                bg = extract_bg_from_styles(parsed_styles, [".compare-card.recommend"], soup)
            else:
                bg = extract_bg_from_styles(parsed_styles, [".compare-card"], soup)
        if is_recommend:
            bg = bg or "e3c877"
            apply_shading_to_cell(cell, bg)
            set_cell_borders(cell, top_color=None, color=border_color)
        else:
            bg = bg or "ffffff"
            apply_shading_to_cell(cell, bg)
            set_cell_borders(cell, top_color=accent_color, color=border_color)

        if is_recommend:
            title_color = extract_color_from_styles(parsed_styles,
                                                    [".compare-card.recommend h3"], soup,
                                                    default="ffffff")
            tag_color = extract_color_from_styles(parsed_styles,
                                                  [".compare-card.recommend .tagline"], soup,
                                                  default="ffffff")
            body_color = extract_color_from_styles(parsed_styles,
                                                   [".compare-card.recommend p"], soup,
                                                   default="ffffff")
        else:
            title_color = extract_color_from_styles(parsed_styles,
                                                    [".compare-card h3"], soup,
                                                    default=palette.get("primary", "1e293b"))
            tag_color = extract_color_from_styles(parsed_styles,
                                                  [".compare-card .tagline"], soup,
                                                  default=palette.get("accent", "a37d14"))
            body_color = extract_color_from_styles(parsed_styles,
                                                   [".compare-card p"], soup,
                                                   default=palette.get("text_secondary", "6b5e50"))

        tagline = card.find(class_="tagline")
        h3 = card.find("h3")
        ps = card.find_all("p")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=60, start=80, bottom=60, end=80)

        first_para = cell.paragraphs[0]
        if tagline:
            run_tag = first_para.add_run(get_text_content(tagline))
            set_run_font(run_tag, size=Pt(10), bold=True, color=hex_to_rgb(tag_color))
        if h3:
            p_title = add_cell_paragraph(cell) if (tagline or first_para.text) else first_para
            p_title.paragraph_format.space_before = Pt(4)
            run_title = p_title.add_run(get_text_content(h3))
            set_run_font(run_title, size=Pt(14), bold=True, color=hex_to_rgb(title_color))
        for p in ps:
            text = get_text_content(p)
            if not text:
                continue
            p_body = add_cell_paragraph(cell)
            p_body.paragraph_format.space_before = Pt(2)
            _add_rich_text_runs(p_body, p, default_size=Pt(10), default_color=body_color)


def process_comparison(doc, section, palette=None):
    """左右对比 → 两列对比表格 + 富文本写入"""
    sides = [c for c in section.children if isinstance(c, Tag) and c.get_text(strip=True)]
    if len(sides) >= 2:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if flush_pending_page_break(doc):
            first_cell = table.cell(0, 0)
            if first_cell.paragraphs:
                inject_page_break_before(first_cell.paragraphs[0])
        text_color = (palette or {}).get("text", "1f2d3d")
        for j, side in enumerate(sides[:2]):
            cell = table.cell(0, j)
            bg = extract_bg_with_gradient_fallback(side.get("style", ""))
            if bg:
                apply_shading_to_cell(cell, bg)
            # 富文本写入：遍历 side 内子元素
            first_para = True
            for child in side.children:
                if not isinstance(child, Tag):
                    continue
                if child.name == "p":
                    p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                    first_para = False
                    _add_rich_text_runs(p_cell, child, default_size=Pt(11), default_color=text_color)
                elif child.name in ("h3", "h4"):
                    p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                    first_para = False
                    p_cell.paragraph_format.space_before = Pt(4)
                    p_cell.paragraph_format.space_after = Pt(2)
                    run = p_cell.add_run(get_text_content(child))
                    set_run_font(run, size=Pt(13), bold=True, color=hex_to_rgb((palette or {}).get("primary", "1e293b")))
                elif child.name in ("ul", "ol"):
                    for li in child.find_all("li", recursive=False):
                        p_cell = add_cell_paragraph(cell)
                        _add_rich_text_runs(p_cell, li, default_size=Pt(11), default_color=text_color)
                else:
                    text = get_text_content(child)
                    if text:
                        p_cell = cell.paragraphs[0] if first_para else add_cell_paragraph(cell)
                        first_para = False
                        _add_rich_text_runs(p_cell, child, default_size=Pt(11), default_color=text_color)
            if first_para:
                cell.text = side.get_text(strip=True)
                for para in cell.paragraphs:
                    for run in para.runs:
                        set_run_font(run, size=Pt(11), color=hex_to_rgb(text_color))
            set_cell_margins(cell, top=30, start=60, bottom=30, end=60)
        return
    # 兜底：表格形式
    table_elem = section.find("table")
    if table_elem:
        process_table(doc, table_elem, palette, None, None)
        return
    flush_pending_page_break(doc)


def process_faq(doc, section, palette):
    """问答对 → 加粗问题 + 缩进段落，保留 <strong>/<span> 等富文本格式

    兼容两类输入：
    1. <details><summary>标题</summary>正文</details> —— 折叠卡片形式。
       注意：section 本身可能就是 <details>（兜底识别），用 find() 找不到自身，
       必须判断 section.name；正文部分改为遍历所有块级子元素
       （p/div/pre/ul/ol 等），避免只取 <p> 导致折叠卡内代码块被吞。
    2. 问答对（div.faq-item / 裸 p 序列）—— 原有逻辑保持不变。
    """
    has_pending = flush_pending_page_break(doc)
    primary_color = palette.get("primary", "1e293b")

    def _emit_summary_title(summary):
        """折叠标题 → 加粗段落"""
        if not summary:
            return
        para = doc.add_paragraph()
        if has_pending:
            inject_page_break_before(para)
        _add_rich_text_runs(para, summary, default_size=Pt(11), default_color=primary_color)
        for run in para.runs:
            run.bold = True

    def _emit_block_children(container):
        """遍历容器所有块级子元素，各自输出为段落。

        覆盖 <p>/<div>/<pre>/<ul>/<ol>/<table>/<h1-h6> 等，未识别的块级元素
        按 get_text_content 兜底为段落——宁可全部输出，不可静默丢弃。
        行内元素（span/strong/em/code 等）不在此层遍历（由父级富文本处理）。
        """
        from html2docx_enhanced import process_code_block as _pcode
        from html2docx_core import process_bullet_list as _pblist
        from html2docx_core import process_table as _ptable
        BLOCK_NAMES = ("p", "div", "section", "article", "pre", "ul", "ol",
                       "h1", "h2", "h3", "h4", "h5", "h6", "table", "figure")
        for child in container.children:
            if not isinstance(child, Tag):
                continue
            if isinstance(child, Comment):
                continue
            if child.name in ("style", "script"):
                continue
            if child.name == "summary":
                continue  # 标题已单独输出，避免重复
            # 跳过交互辅助元素（复制按钮/搜索框/返回顶部等），避免误输出按钮文本
            if any(c in ("code-copy-btn", "table-search-input", "table-no-result",
                         "back-to-top", "dark-toggle-btn")
                   for c in (child.get("class") or [])):
                continue
            if child.name == "pre":
                _pcode(doc, child)
                continue
            if child.name in ("ul", "ol"):
                _pblist(doc, child)
                continue
            if child.name == "table":
                _ptable(doc, child, palette, None, None)
                continue
            if child.name in BLOCK_NAMES:
                if child.name in ("div", "section", "article"):
                    _emit_block_children(child)
                else:
                    para = doc.add_paragraph()
                    if has_pending:
                        inject_page_break_before(para)
                    para.paragraph_format.left_indent = Cm(0.5)
                    _add_rich_text_runs(para, child, default_size=Pt(11), default_color="1f2d3d")
                continue
            # 其余块级/未知元素：按可见文本兜底输出（宁丑勿丢）
            text = get_text_content(child)
            if text:
                _add_rich_text_runs(doc.add_paragraph(), child,
                                    default_size=Pt(11), default_color="1f2d3d")
        # 容器内若只有裸文本（无块级子元素），兜底整体输出
        bare_text = "".join(str(c) for c in container.children
                            if isinstance(c, NavigableString) and not isinstance(c, Comment)).strip()
        if bare_text:
            # 裸文本是 str 类型，_add_rich_text_runs 只认 Tag/NavigableString，
            # 必须包装为 NavigableString，否则段落建了但一个字不写（如 .code-lang-label）
            _add_rich_text_runs(doc.add_paragraph(), NavigableString(bare_text),
                                default_size=Pt(11), default_color="1f2d3d")

    def _add_block_para(container, indent=False):
        """为容器整体创建段落（避免与标题混在同一段）"""
        para = doc.add_paragraph()
        if has_pending:
            inject_page_break_before(para)
        if indent:
            para.paragraph_format.left_indent = Cm(0.5)
        _add_rich_text_runs(para, container, default_size=Pt(11), default_color="1f2d3d")

    # details/summary 形式：section 本身可能是 details，也可能是其父容器
    details = section if section.name == "details" else section.find("details")
    if details:
        summary = details.find("summary")
        _emit_summary_title(summary)
        _emit_block_children(details)
        return
    divs = section.find_all("div", recursive=False)
    if not divs:
        ps = section.find_all("p")
        first_p = True
        for p in ps:
            text = get_text_content(p)
            if not text:
                continue
            if text.startswith("Q:") or text.startswith("问"):
                para = doc.add_paragraph()
                if first_p and has_pending:
                    inject_page_break_before(para)
                    first_p = False
                _add_rich_text_runs(para, p, default_size=Pt(11), default_color=primary_color)
                for run in para.runs:
                    run.bold = True
            else:
                para = doc.add_paragraph()
                if first_p and has_pending:
                    inject_page_break_before(para)
                    first_p = False
                para.paragraph_format.left_indent = Cm(0.5)
                _add_rich_text_runs(para, p, default_size=Pt(11), default_color="1f2d3d")
        return
    for div in divs:
        ps = div.find_all("p")
        first_p_in_div = True
        for p in ps:
            text = get_text_content(p)
            if not text:
                continue
            if text.startswith("Q:") or text.startswith("问"):
                para = doc.add_paragraph()
                if first_p_in_div and has_pending:
                    inject_page_break_before(para)
                    first_p_in_div = False
                _add_rich_text_runs(para, p, default_size=Pt(11), default_color=primary_color)
                for run in para.runs:
                    run.bold = True
            else:
                para = doc.add_paragraph()
                if first_p_in_div and has_pending:
                    inject_page_break_before(para)
                    first_p_in_div = False
                para.paragraph_format.left_indent = Cm(0.5)
                _add_rich_text_runs(para, p, default_size=Pt(11), default_color="1f2d3d")


def process_badge_group(doc, section, palette):
    """标签组 → 逗号分隔文本"""
    spans = section.find_all("span")
    texts = [get_text_content(s) for s in spans if get_text_content(s)]
    if not texts:
        # 兜底：整段文本按顿号拆分
        whole = get_text_content(section)
        texts = [t for t in re.split(r"[、,，]", whole) if t.strip()]
    if texts:
        para = doc.add_paragraph("、".join(texts))
        if flush_pending_page_break(doc):
            inject_page_break_before(para)
        for run in para.runs:
            set_run_font(run, size=Pt(11), color=hex_to_rgb(palette.get("accent", "6366f1")))
    else:
        flush_pending_page_break(doc)


def process_gallery(doc, section):
    """图片画廊 → 图片纵列"""
    has_pending = flush_pending_page_break(doc)
    imgs = section.find_all("img")
    first_img = True
    for img in imgs:
        img_bytes = _extract_img_bytes(img)
        if img_bytes:
            try:
                if first_img and has_pending:
                    pb_para = doc.add_paragraph()
                    inject_page_break_before(pb_para)
                    pb_para.paragraph_format.space_before = Pt(0)
                    pb_para.paragraph_format.space_after = Pt(0)
                    pb_para.paragraph_format.line_spacing = Pt(1)
                    first_img = False
                doc.add_picture(io.BytesIO(img_bytes), width=Inches(4.5))
                # 图片所在段落居中
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                print(f"[WARN] 画廊图片处理失败: {e}")
    # 图注
    for img in imgs:
        figcaption = img.find_parent("figure")
        if figcaption:
            cap = figcaption.find("figcaption")
            if cap and get_text_content(cap):
                para = doc.add_paragraph(get_text_content(cap))
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run_font(run, size=Pt(10), color=hex_to_rgb("64748b"))


def process_pricing_table(doc, section, palette=None):
    """定价/方案对比表 → 对比表格（shading 高亮列）"""
    table_elem = section.find("table") if section.name != "table" else section
    if not table_elem:
        flush_pending_page_break(doc)
        return
    rows, merges, _detail = parse_table_data(table_elem)
    if not rows:
        flush_pending_page_break(doc)
        return
    # 高亮列：thead 中带非纯白/非页面底色 th 所在列
    # 排除规则：纯白（ffffff）、当前调色板页面底色（palette.bg），其余视为高亮列
    highlight_cols = set()
    palette_bg = ((palette or {}).get("bg") or "f8fafc").lstrip("#").lower()
    exclude_colors = {"ffffff", palette_bg}
    thead = table_elem.find("thead")
    if thead:
        for th_idx, th in enumerate(thead.find_all("th")):
            th_bg = extract_bg_with_gradient_fallback(th.get("style", ""))
            if th_bg and th_bg.lower() not in exclude_colors:
                highlight_cols.add(th_idx)
    # 表头判定：<thead> 存在或首行含 <th> 标签（原实现引用未定义的 has_header，触发 NameError）
    # 跳过 .row-detail 行：与 parse_table_data 的过滤保持一致
    tr_tags = [tr for tr in table_elem.find_all("tr")
               if "row-detail" not in (tr.get("class") or [])]
    first_row_has_th = bool(tr_tags and tr_tags[0].find(["th"], recursive=False))
    has_header = thead is not None or first_row_has_th
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if flush_pending_page_break(doc):
        first_cell = table.cell(0, 0)
        if first_cell.paragraphs:
            inject_page_break_before(first_cell.paragraphs[0])
    # 高亮底色用调色板 accent_light（如无则用浅青 f0fdfa 兜底）
    highlight_fill = (palette or {}).get("accent_light", "f0fdfa")
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.cell(i, j)
                # 富文本写入：遍历 HTML td/th 子节点
                td_el = None
                if i < len(tr_tags):
                    cells_in_row = tr_tags[i].find_all(["td", "th"], recursive=False)
                    if j < len(cells_in_row):
                        td_el = cells_in_row[j]
                if td_el is not None:
                    para = cell.paragraphs[0]
                    is_header = has_header and i == 0
                    d_size = Pt(10.5) if is_header else Pt(11)
                    d_color = (palette or {}).get("primary", "1e293b") if is_header else (palette or {}).get("text", "1f2d3d")
                    _add_rich_text_runs(para, td_el, default_size=d_size, default_color=d_color)
                    if is_header:
                        for run in para.runs:
                            run.bold = True
                else:
                    cell.text = cell_text
                    if has_header and i == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                                set_run_font(run, size=Pt(10.5))
                    else:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_run_font(run, size=Pt(11))
                set_cell_margins(cell, top=30, start=60, bottom=30, end=60)
                if j in highlight_cols and i > 0:
                    apply_shading_to_cell(cell, highlight_fill)

    # 重建合并单元格（colspan/rowspan），从后往前避免索引错位
    for m in sorted(merges, key=lambda x: (-x["r"], -x["c"])):
        r, c, rs, cs = m["r"], m["c"], m["rs"], m["cs"]
        if r >= len(rows) or c >= num_cols:
            continue
        r_end = min(r + rs - 1, len(rows) - 1)
        c_end = min(c + cs - 1, num_cols - 1)
        if r_end == r and c_end == c:
            continue
        try:
            table.cell(r, c).merge(table.cell(r_end, c_end))
        except Exception:
            pass


def process_icon_list(doc, section, palette):
    """带图标列表 → 普通列表（图标丢失），保留 <strong>/<span> 等富文本格式"""
    has_pending = flush_pending_page_break(doc)
    lis = section.find_all("li")
    if not lis:
        # 逐行文本
        for p in section.find_all("p"):
            text = get_text_content(p)
            if text:
                para = doc.add_paragraph()
                if has_pending:
                    inject_page_break_before(para)
                    has_pending = False
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                _add_rich_text_runs(para, p, default_size=Pt(11))
        return
    first = True
    for li in lis:
        text = get_text_content(li)
        if not text:
            continue
        para = doc.add_paragraph()
        if first and has_pending:
            inject_page_break_before(para)
            first = False
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        _add_rich_text_runs(para, li, default_size=Pt(11))


def process_end_page(doc, section, palette, parsed_styles=None, soup=None):
    """尾页/致谢页 → 分页 + 1x1 表格垂直居中 + 背景色 shading

    遍历 section 所有直接子元素，按类型分别输出（h1/h2/h3 → 标题段落；
    p → 富文本段落，读取内联 font-size/color；div/span → 递归提取文本）。
    不再只提取 h1 + p，避免遗漏 div/span 等标签承载的内容。
    """
    from docx_utils import get_inline_style_prop
    bg_color = extract_bg_with_gradient_fallback(section.get("style", ""))
    if not bg_color:
        # 选择器优先级：.slide.hero（mode-deck 深色渐变尾页）> .end_page > section.end_page
        # 补充 .title.end_page（封面样式复用为尾页，如 class="title end_page" 的致谢页）
        # 及 .title 兜底（尾页背景定义在 .title 规则上时也可取到，如本报告尾页复用封面样式）
        bg_color = extract_bg_from_styles(
            parsed_styles, [".slide.hero", ".end_page", "section.end_page", ".title", "section.title", ".title.end_page", "section.title.end_page"], soup)

    # 1x1 表格承载尾页，实现上下左右居中（与封面一致）
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_nil(table)
    set_table_width_percent(table, 100)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
    # 行高动态计算：按页面实际可用高度设置（与封面一致），
    # 避免写死 21.5cm 导致尾页内容偏上。
    _page_h = 29.7
    try:
        _sec = doc.sections[0]
        _page_h = _sec.page_height / 360000.0  # EMU → cm
    except Exception:
        pass
    _top_m = 1.5
    _bot_m = 1.5
    try:
        _top_m = doc.sections[0].top_margin / 360000.0
        _bot_m = doc.sections[0].bottom_margin / 360000.0
    except Exception:
        pass
    _avail_h = max(_page_h - _top_m - _bot_m, 20.0)
    # 尾页表格行高减 4cm：给表格后的 AIGC 水印段落（2 段，默认行距+段后间距）
    # 留出足够空间，避免水印被挤到下一页形成空白页。
    _end_row_h = max(_avail_h - 4.0, 20.0)
    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row.height = Cm(_end_row_h)

    # 尾页文字颜色：优先读取 HTML <style> 中 .slide.hero h1 / .end_page h1 的 color 定义
    end_h1_color = extract_color_from_styles(
        parsed_styles,
        [".title.end_page h1", ".slide.hero h1", ".end_page h1", "section.end_page h1", ".end_page > h1", ".title h1"],
        soup,
        default=palette.get("primary", "1e293b"),
    )
    end_p_color = extract_color_from_styles(
        parsed_styles,
        [".title.end_page p", ".slide.hero p", ".end_page p", "section.end_page p", ".end_page > p", ".title p"],
        soup,
        default=palette.get("accent", "64748b"),
    )

    first_content = True
    for child in section.children:
        if not isinstance(child, Tag):
            continue
        # 跳过纯装饰元素（accent-line 等）
        child_classes = child.get("class", [])
        if "accent-line" in child_classes:
            continue

        if child.name in ("h1", "h2", "h3"):
            # 标题段落
            para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
            if first_content and flush_pending_page_break(doc):
                inject_page_break_before(para)
            first_content = False
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(8)
            if bg_color:
                apply_shading(para, bg_color)
            run = para.add_run(get_text_content(child))
            set_run_font(run, size=Pt(28), bold=True, color=hex_to_rgb(end_h1_color))
        elif child.name == "p":
            # 富文本段落：读取内联 font-size/color，回退 .end_page p 样式
            para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
            if first_content and flush_pending_page_break(doc):
                inject_page_break_before(para)
            first_content = False
            if bg_color:
                apply_shading(para, bg_color)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            inline_size = get_inline_style_prop(child.get("style", ""), "font-size")
            inline_color = get_inline_style_prop(child.get("style", ""), "color")
            d_size = Pt(inline_size) if inline_size else Pt(14)
            d_color = inline_color or end_p_color
            _add_rich_text_runs(para, child, default_size=d_size, default_color=d_color)
        elif child.name == "div":
            child_classes = child.get("class", [])
            # 检测 .title-meta 容器：提取 .meta-item 列表，用嵌套表格横排（与封面一致）
            if "title-meta" in child_classes:
                meta_items = []
                for item in child.find_all(class_="meta-item"):
                    val_el = item.find(class_="value")
                    label_el = item.find(class_="label")
                    if val_el and label_el:
                        meta_items.append((get_text_content(val_el), get_text_content(label_el)))
                if meta_items:
                    # 分隔线
                    sep_para = add_cell_paragraph(cell) if not first_content else cell.paragraphs[0]
                    if first_content and flush_pending_page_break(doc):
                        inject_page_break_before(sep_para)
                    first_content = False
                    sep_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    sep_para.paragraph_format.space_before = Pt(8)
                    sep_para.paragraph_format.space_after = Pt(4)
                    if bg_color:
                        apply_shading(sep_para, bg_color)
                    sep_run = sep_para.add_run("─" * 8)
                    set_run_font(sep_run, size=Pt(8), color=hex_to_rgb(palette.get("accent", "6366f1")))
                    # 嵌套表格：每列一个 meta-item
                    meta_table = cell.add_table(rows=1, cols=len(meta_items))
                    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders_nil(meta_table)
                    set_table_width_percent(meta_table, 100)
                    for j, (val, label) in enumerate(meta_items):
                        m_cell = meta_table.cell(0, j)
                        m_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if bg_color:
                            apply_shading_to_cell(m_cell, bg_color)
                        # 数值（大字号加粗）
                        p_val = m_cell.paragraphs[0]
                        p_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_val.paragraph_format.space_before = Pt(2)
                        p_val.paragraph_format.space_after = Pt(2)
                        run_val = p_val.add_run(val)
                        set_run_font(run_val, size=Pt(15), bold=True,
                                     color=hex_to_rgb("#ffffff" if is_dark_color(bg_color) else end_h1_color))
                        # 标签（小字号浅色）
                        p_label = m_cell.add_paragraph()
                        p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_label.paragraph_format.space_before = Pt(0)
                        p_label.paragraph_format.space_after = Pt(2)
                        run_label = p_label.add_run(label)
                        set_run_font(run_label, size=Pt(9), color=hex_to_rgb(end_p_color))
                    continue  # title-meta 已处理，跳过通用 div 递归
            # 其他 div 容器：递归提取内部子元素
            # 修复（2026-08-28）：纯文本 div（如 <div>THANK YOU</div>，只有文本无子标签）
            # 此前被 for 循环中的 `if not isinstance(sub, Tag): continue` 整体跳过 → 内容丢失。
            # 现在：div 无 Tag 子元素但有可见文本 → 按段落输出（读取内联 font-size/color）。
            tag_subs = [s for s in child.children if isinstance(s, Tag)]
            div_text = get_text_content(child)
            if not tag_subs and div_text:
                para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                if first_content and flush_pending_page_break(doc):
                    inject_page_break_before(para)
                first_content = False
                if bg_color:
                    apply_shading(para, bg_color)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                inline_size = get_inline_style_prop(child.get("style", ""), "font-size")
                inline_color = get_inline_style_prop(child.get("style", ""), "color")
                d_size = Pt(inline_size) if inline_size else Pt(14)
                d_color = inline_color or end_p_color
                _add_rich_text_runs(para, child, default_size=d_size, default_color=d_color)
            for sub in child.children:
                if not isinstance(sub, Tag):
                    continue
                if sub.name in ("h1", "h2", "h3"):
                    para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                    if first_content and flush_pending_page_break(doc):
                        inject_page_break_before(para)
                    first_content = False
                    if bg_color:
                        apply_shading(para, bg_color)
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run(get_text_content(sub))
                    set_run_font(run, size=Pt(28), bold=True, color=hex_to_rgb(end_h1_color))
                elif sub.name == "p" or get_text_content(sub):
                    para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                    if first_content and flush_pending_page_break(doc):
                        inject_page_break_before(para)
                    first_content = False
                    if bg_color:
                        apply_shading(para, bg_color)
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_before = Pt(2)
                    para.paragraph_format.space_after = Pt(2)
                    inline_size = get_inline_style_prop(sub.get("style", ""), "font-size")
                    inline_color = get_inline_style_prop(sub.get("style", ""), "color")
                    d_size = Pt(inline_size) if inline_size else Pt(14)
                    d_color = inline_color or end_p_color
                    _add_rich_text_runs(para, sub, default_size=d_size, default_color=d_color)
        elif child.name == "span":
            # span 标签：提取文本输出
            text = get_text_content(child)
            if text:
                para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                if first_content and flush_pending_page_break(doc):
                    inject_page_break_before(para)
                first_content = False
                if bg_color:
                    apply_shading(para, bg_color)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                inline_size = get_inline_style_prop(child.get("style", ""), "font-size")
                inline_color = get_inline_style_prop(child.get("style", ""), "color")
                d_size = Pt(inline_size) if inline_size else Pt(14)
                d_color = inline_color or end_p_color
                _add_rich_text_runs(para, child, default_size=d_size, default_color=d_color)
        elif child.name == "a":
            # a.cta-btn 按钮（立即购买/了解更多等）：独立段落 + 加粗高亮
            text = get_text_content(child)
            if text:
                para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                if first_content and flush_pending_page_break(doc):
                    inject_page_break_before(para)
                first_content = False
                if bg_color:
                    apply_shading(para, bg_color)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(4)
                run = para.add_run(text)
                child_classes = child.get("class", []) or []
                if "cta-btn" in child_classes:
                    # 主按钮：白字 + accent 底色
                    set_run_font(run, size=Pt(11), bold=True,
                                 color=hex_to_rgb("ffffff"))
                    apply_shading_to_run(run, palette.get("accent", "3b82f6"))
                else:
                    set_run_font(run, size=Pt(11), bold=True,
                                 color=hex_to_rgb(end_h1_color))
        elif child.name == "img":
            # 图片（base64）：插入尾页单元格，避免图片丢失
            try:
                img_bytes = _extract_img_bytes(child)
                if img_bytes:
                    # 通过 doc.add_picture 注册到文档 part，再把 shape 移入单元格
                    # （add_cell_paragraph 的 parent 是 CT_Tc 无 part，直接 add_picture 会报错）
                    from copy import deepcopy
                    inline_shape = doc.add_picture(io.BytesIO(img_bytes), width=Inches(3.5))
                    temp_para = doc.paragraphs[-1]
                    shape_elem = deepcopy(inline_shape._inline)
                    temp_para._element.getparent().remove(temp_para._element)
                    p_img = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                    if first_content and flush_pending_page_break(doc):
                        inject_page_break_before(p_img)
                    first_content = False
                    if bg_color:
                        apply_shading(p_img, bg_color)
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img._element.append(shape_elem)
            except Exception:
                pass
        else:
            # 兜底：有文本则输出为段落
            text = get_text_content(child)
            if text:
                para = cell.paragraphs[0] if first_content else add_cell_paragraph(cell)
                if first_content and flush_pending_page_break(doc):
                    inject_page_break_before(para)
                first_content = False
                if bg_color:
                    apply_shading(para, bg_color)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                run = para.add_run(text)
                set_run_font(run, size=Pt(14), color=hex_to_rgb(end_p_color))
